from Security.Advance_Logger import AdvancedLogger
from pathlib import Path
from email import policy
from email.parser import BytesParser
from bs4 import BeautifulSoup
import pymupdf4llm
import openpyxl
import docx
import re

logger: AdvancedLogger = AdvancedLogger() 

class FileParser:
    def sanitize_text(user_input: str) -> str:
        """
        Removes potentially dangerous characters and trims whitespace.
        """
        if not isinstance(user_input, str):
            raise ValueError("Input must be a string")
        
        safe_text = re.sub(r"[^a-zA-Z0-9\s.,!?@#_-]", "", user_input)
        return safe_text.strip()

    @staticmethod
    def parse_pdf(path: str) -> str:
        try:
            markdown = pymupdf4llm.to_markdown(path)
            return markdown
        except Exception as e:
            logger.error("FileParser.parse_pdf", e)
            return ""
    
    @staticmethod
    def parse_docx(path: str) -> str:
        try:
            doc = docx.Document(path)
            full_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            
            return full_text if full_text else ""
            
        except Exception as e:
            logger.error("FileParser.parse_docx", e)
            return "" 

    @staticmethod
    def extract_codebase(path: str) -> str:
        with open(path, "rb") as f:
            return f.read().decode("utf-8", errors="ignore")
    
    @staticmethod
    def extract_email(path: str) -> str:

        with open(path, "rb") as f:
            msg = BytesParser(
                policy=policy.default
            ).parse(f)

        body = None

        # Prefer plain text immediately
        for part in msg.walk():

            ctype = part.get_content_type()

            if ctype == "text/plain":

                try:
                    body = part.get_content()
                    break

                except Exception:
                    pass

        # Fallback to HTML
        if body is None:

            for part in msg.walk():

                if part.get_content_type() == "text/html":

                    try:

                        html = part.get_content()

                        soup = BeautifulSoup(
                            html,
                            "lxml"
                        )

                        for tag in soup([
                            "script",
                            "style",
                            "img",
                            "meta",
                            "head"
                        ]):
                            tag.decompose()

                        body = soup.get_text("\n")
                        break

                    except Exception:
                        pass

        if not body:
            body = ""

        lines = []

        append = lines.append

        for line in body.splitlines():

            line = line.strip()

            if not line:
                continue

            lower = line.lower()

            if (
                "unsubscribe" in lower
                or "tracking" in lower
                or "emailopen" in lower
            ):
                continue

            append(line)

        return "\n".join(lines)

    @staticmethod
    def parse_excel(path: str) -> str:
        try:
            rows = []
            append = rows.append
            
            if path.lower().endswith(".csv"):
                import csv

                with open(
                    path,
                    "r",
                    encoding="utf-8",
                    errors="ignore",
                    newline=""
                ) as f:

                    reader = csv.reader(f)

                    for row in reader:

                        values = [
                            str(cell).strip()
                            for cell in row
                            if cell not in (None, "")
                        ]

                        if values:
                            append(" | ".join(values))

                return "\n".join(rows).strip()
            wb = openpyxl.load_workbook(
                path,
                data_only=True,
                read_only=True
            )

            rows = []
            append = rows.append

            for sheet in wb.worksheets:

                append(f"\nSheet: {sheet.title}")

                for row in sheet.iter_rows(values_only=True):

                    values = [
                        str(cell).strip()
                        for cell in row
                        if cell is not None
                    ]

                    if values:
                        append(" | ".join(values))

            return "\n".join(rows).strip()

        except Exception as e:
            logger.error("FileParser.parse_excel", e)
            return ""
        
class ParseFile:

    PDF = ".pdf"
    DOCX = ".docx"
    EML = ".eml"

    XL = frozenset({
        ".xlsx", ".xls", ".csv"
    })

    CODE_EXTENSIONS = frozenset({
    ".py", ".pyi",
    ".js", ".jsx", ".mjs", ".cjs",
    ".ts", ".tsx",
    ".java",
    ".c", ".h",
    ".cpp", ".cc", ".cxx", ".hpp", ".hh", ".hxx",
    ".cs",
    ".go",
    ".rs",
    ".php",
    ".rb",
    ".swift",
    ".kt", ".kts",
    ".scala",
    ".dart",
    ".r",
    ".lua",
    ".pl",
    ".sh", ".bash", ".zsh",
    ".ps1",
    ".sql",
    ".html", ".htm",
    ".css", ".scss", ".sass", ".less",
    ".json",
    ".xml",
    ".yaml", ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".env",
    ".md",
    ".txt",
    ".dockerfile",
    ".tf",
    ".vue",
    ".svelte",
    ".ipynb"
})

    @staticmethod
    def parse_any_file(path: str) -> str:

        try:

            suffix = Path(path).suffix.lower()

            # PDF
            if suffix == ParseFile.PDF:
                return FileParser.parse_pdf(path)

            # DOCX
            if suffix == ParseFile.DOCX:
                return FileParser.parse_docx(path)

            # EMAIL
            if suffix == ParseFile.EML:
                return FileParser.extract_email(path)

            # CODE / TEXT
            if suffix in ParseFile.CODE_EXTENSIONS:
                return FileParser.extract_codebase(path)
            
            if suffix in ParseFile.XL:
                return FileParser.parse_excel(path)

            # FALLBACK
            return FileParser.extract_codebase(path)

        except Exception as e:
            logger.error("ParseFile.parse_any_file", e)
            return ""
        
class Chunker:
    @staticmethod
    def chunk_text(text: str, chunk_size: int = 1200, chunk_overlap: int = 200) -> list[str]:
        """
        Splits standard text files into overlapping segments to keep context intact.
        """
        if not text or not text.strip():
            return []
        
        chunks = []
        start = 0
        text_len = len(text)
        
        while start < text_len:
            end = start + chunk_size
            chunk = text[start:end]
            chunks.append(chunk.strip())
            # Shift forward by the size minus overlap to create a sliding window
            start += (chunk_size - chunk_overlap)
            
        return chunks

    @staticmethod
    def chunk_code(code: str, language_suffix: str) -> list[str]:
        """
        Heuristic splitter designed to isolate logical blocks in programming code.
        """
        if not code or not code.strip():
            return []

        # Split by lines to analyze block structure
        lines = code.splitlines()
        chunks = []
        current_chunk = []
        current_length = 0
        
        brace_languages = {".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".h", ".cpp", ".cc", ".cxx", ".hpp", ".hh", ".hxx", ".cs", ".go", ".rs", ".php", ".swift", ".kt", ".kts", ".scala", ".dart", ".groovy", ".m", ".mm", ".zig", ".vue", ".css", ".scss", ".sass", ".less", ".json", }
        
        brace_count = 0
        
        for line in lines:
            current_chunk.append(line)
            current_length += len(line)
            
            if language_suffix in brace_languages:
                brace_count += line.count("{") - line.count("}")
                if brace_count == 0 and current_length > 800:
                    chunks.append("\n".join(current_chunk))
                    current_chunk = []
                    current_length = 0
            else:
                if (line.startswith("def ") or line.startswith("class ") or line.startswith("import ")) and current_length > 800:
                    next_start = current_chunk.pop()
                    if current_chunk:
                        chunks.append("\n".join(current_chunk))
                    current_chunk = [next_start]
                    current_length = len(next_start)

        if current_chunk:
            chunks.append("\n".join(current_chunk))
            
        return chunks
    
if __name__ == "__main__":
    lis = [r"F:\Smart AI workflow\Workflow\Frontend_Connection.py"]
    for i in lis:
        print(Chunker.chunk_code(ParseFile.parse_any_file(i), ".py"))